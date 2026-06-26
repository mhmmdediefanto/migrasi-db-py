from __future__ import annotations

import csv
import os
from datetime import date
from pathlib import Path

from lib.db import connect_pg, fetch_one
from lib.progress import Spinner

ROOT = Path(__file__).resolve().parent.parent
OUTPUT_DIR = ROOT / "output"

DEFAULT_SINCE_DATE = date(2026, 6, 18)

OUTPUT_FIELDS = [
    "kode_barang",
    "nama_barang",
    "stok_excel",
    "stok_barang_system",
    "selisih_barang",
    "ada_pembelian",
    "pembelian_tanggal_terakhir",
    "pembelian_jumlah_faktur",
    "pembelian_setelah_laporan",
    "ada_penjualan",
    "penjualan_tanggal_terakhir",
    "penjualan_jumlah_faktur",
    "penjualan_setelah_laporan",
]

PEMBELIAN_SQL = """
    SELECT
        b.kode_barang,
        COUNT(DISTINCT p.id) AS jumlah_faktur,
        MAX(p.tanggal) AS tanggal_terakhir,
        COUNT(DISTINCT p.id) FILTER (WHERE p.tanggal >= %s) AS setelah_laporan
    FROM barang b
    INNER JOIN pembelian_detail pd ON pd.barang_id = b.id AND pd.deleted_at IS NULL
    INNER JOIN pembelian p ON p.id = pd.pembelian_id AND p.deleted_at IS NULL
    WHERE b.deleted_at IS NULL
      AND b.kode_barang = ANY(%s)
      AND p.gudang_id = %s
    GROUP BY b.kode_barang
"""

PENJUALAN_SQL = """
    SELECT
        b.kode_barang,
        COUNT(DISTINCT p.id) AS jumlah_faktur,
        MAX(p.tanggal::date) AS tanggal_terakhir,
        COUNT(DISTINCT p.id) FILTER (WHERE p.tanggal::date >= %s) AS setelah_laporan
    FROM barang b
    INNER JOIN penjualan_detail pd ON pd.barang_id = b.id AND pd.deleted_at IS NULL
    INNER JOIN penjualan p ON p.id = pd.penjualan_id AND p.deleted_at IS NULL
    WHERE b.deleted_at IS NULL
      AND b.kode_barang = ANY(%s)
    GROUP BY b.kode_barang
"""


def _ensure_output_dir() -> Path:
    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
    return OUTPUT_DIR


def load_selisih_rows(path: Path) -> list[dict]:
    from openpyxl import load_workbook

    workbook = load_workbook(path, read_only=True, data_only=True)
    sheet = workbook.active
    headers = [cell.value for cell in next(sheet.iter_rows(min_row=1, max_row=1))]
    field_map = {str(h): i for i, h in enumerate(headers) if h}

    def col(name: str, fallback: int) -> int:
        return field_map.get(name, fallback)

    rows: list[dict] = []
    for row in sheet.iter_rows(min_row=2, values_only=True):
        if not row or not any(row):
            continue
        kode = str(row[col("kode_barang", 0)] or "").strip()
        nama = str(row[col("nama_barang", 1)] or "").strip()
        if not kode:
            continue
        rows.append(
            {
                "kode_barang": kode,
                "nama_barang": nama,
                "stok_excel": row[col("stok_excel", 2)],
                "stok_barang_system": row[col("stok_barang_system", 3)],
                "selisih_barang": row[col("selisih_barang", 4)],
            }
        )
    workbook.close()
    return rows


def resolve_gudang_id(pg, gudang_hint: str) -> tuple[int, str]:
    from lib.stok_check import GUDANG_ALIASES

    key = gudang_hint.strip().lower()
    name = GUDANG_ALIASES.get(key, gudang_hint.strip())
    gudang_id = fetch_one(
        pg,
        "SELECT id FROM master_gudang WHERE name ILIKE %s ORDER BY id LIMIT 1",
        (name,),
    )
    if gudang_id is None:
        raise RuntimeError(f"Gudang tidak ditemukan: {gudang_hint!r}")
    return int(gudang_id), name


def _fetch_pembelian_map(pg, since: date, kodes: list[str], gudang_id: int) -> dict[str, dict]:
    with pg.cursor() as cur:
        cur.execute(PEMBELIAN_SQL, (since, kodes, gudang_id))
        rows = cur.fetchall()
    return _rows_to_map(rows)


def _fetch_penjualan_map(pg, since: date, kodes: list[str]) -> dict[str, dict]:
    with pg.cursor() as cur:
        cur.execute(PENJUALAN_SQL, (since, kodes))
        rows = cur.fetchall()
    return _rows_to_map(rows)


def _rows_to_map(rows) -> dict[str, dict]:
    result: dict[str, dict] = {}
    for kode, jumlah, tanggal_terakhir, setelah in rows:
        result[str(kode)] = {
            "jumlah_faktur": int(jumlah or 0),
            "tanggal_terakhir": tanggal_terakhir,
            "setelah_laporan": int(setelah or 0),
        }
    return result


def _write_xlsx(path: Path, rows: list[dict]) -> None:
    from openpyxl import Workbook

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Trx Selisih"
    sheet.append(OUTPUT_FIELDS)
    for row in rows:
        sheet.append([row.get(field, "") for field in OUTPUT_FIELDS])
    workbook.save(path)


def _write_csv(path: Path, rows: list[dict]) -> None:
    with path.open("w", encoding="utf-8", newline="") as handle:
        writer = csv.DictWriter(handle, fieldnames=OUTPUT_FIELDS, extrasaction="ignore")
        writer.writeheader()
        writer.writerows(rows)


def run_stok_trx_check(
    pg,
    *,
    selisih_file: Path,
    gudang: str = "pusat",
    out_path: Path | None = None,
    since_date: date | None = None,
    only_with_trx: bool = True,
) -> dict:
    path_in = selisih_file.expanduser().resolve()
    if not path_in.is_file():
        raise FileNotFoundError(path_in)

    since = since_date or DEFAULT_SINCE_DATE
    _ensure_output_dir()
    path_out = out_path or OUTPUT_DIR / "stok_selisih_trx_pembelian_penjualan.xlsx"
    if path_out.suffix.lower() != ".xlsx":
        path_out = path_out.with_suffix(".xlsx")

    with Spinner(f"Membaca {path_in.name}..."):
        selisih_rows = load_selisih_rows(path_in)

    kodes = list({r["kode_barang"] for r in selisih_rows})
    gudang_id, gudang_name = resolve_gudang_id(pg, gudang)

    with Spinner(f"Query pembelian ({len(kodes)} kode, 1x)..."):
        pembelian = _fetch_pembelian_map(pg, since, kodes, gudang_id)

    with Spinner(f"Query penjualan ({len(kodes)} kode, 1x)..."):
        penjualan = _fetch_penjualan_map(pg, since, kodes)

    out_rows: list[dict] = []
    with_pembelian = 0
    with_penjualan = 0
    trx_setelah_laporan = 0

    for src in selisih_rows:
        kode = src["kode_barang"]
        p = pembelian.get(kode)
        s = penjualan.get(kode)
        ada_p = p is not None
        ada_s = s is not None
        if ada_p:
            with_pembelian += 1
        if ada_s:
            with_penjualan += 1
        p_after = p.get("setelah_laporan", 0) if p else 0
        s_after = s.get("setelah_laporan", 0) if s else 0
        if p_after > 0 or s_after > 0:
            trx_setelah_laporan += 1

        if only_with_trx and p_after == 0 and s_after == 0:
            continue

        out_rows.append(
            {
                **src,
                "ada_pembelian": "Y" if ada_p else "N",
                "pembelian_tanggal_terakhir": p.get("tanggal_terakhir") if p else "",
                "pembelian_jumlah_faktur": p.get("jumlah_faktur", 0) if p else 0,
                "pembelian_setelah_laporan": p_after,
                "ada_penjualan": "Y" if ada_s else "N",
                "penjualan_tanggal_terakhir": s.get("tanggal_terakhir") if s else "",
                "penjualan_jumlah_faktur": s.get("jumlah_faktur", 0) if s else 0,
                "penjualan_setelah_laporan": s_after,
            }
        )

    try:
        _write_xlsx(path_out, out_rows)
    except ImportError:
        path_out = path_out.with_suffix(".csv")
        _write_csv(path_out, out_rows)

    return {
        "selisih_rows": len(selisih_rows),
        "unique_kode": len(kodes),
        "exported_rows": len(out_rows),
        "only_with_trx": only_with_trx,
        "since_date": since.isoformat(),
        "with_pembelian": with_pembelian,
        "with_penjualan": with_penjualan,
        "trx_setelah_laporan": trx_setelah_laporan,
        "gudang": gudang_name,
        "queries": 2,
        "output_path": str(path_out),
    }


def connect_pg_from_params(
    host: str,
    port: int,
    database: str,
    user: str,
    password: str,
    sslmode: str | None = None,
):
    if sslmode is None:
        explicit = os.getenv("PG_SSLMODE", "").strip()
        if explicit:
            sslmode = explicit
        else:
            sslmode = "require" if host not in ("localhost", "127.0.0.1") else None
    return connect_pg(
        {
            "host": host,
            "port": port,
            "database": database,
            "user": user,
            "password": password,
            "sslmode": sslmode,
        }
    )
