from __future__ import annotations

import re
from datetime import date, datetime
from pathlib import Path

from lib.config import slug_code
from lib.db import fetch_all_dict
from lib.progress import Spinner

ROOT = Path(__file__).resolve().parent.parent
OUTPUT_DIR = ROOT / "output"

DEFAULT_GOLONGAN = "07/MULTIGARMENTAMA + PPN"


def default_output_path(golongan: str, db_name: str = "") -> Path:
    """Nama file otomatis dari golongan + database, mis. stok_minus_07_multigarmentama_ppn_sragen.xlsx."""
    gol_slug = slug_code(golongan, max_length=60).lower()
    db_slug = re.sub(r"[^a-z0-9]+", "_", (db_name or "export").strip().lower()).strip("_")
    return OUTPUT_DIR / f"stok_minus_{gol_slug}_{db_slug}.xlsx"

BARANG_KODE_CTE = """
    barang_kode AS (
        SELECT
            s.cSTKpk,
            MIN(TRIM(sd.cSTDcode)) AS kode_barang,
            TRIM(s.cSTKdesc) AS nama_barang,
            TRIM(g.cGRPdesc) AS golongan
        FROM stock s
        INNER JOIN stockdetail sd ON sd.cSTDfkSTK = s.cSTKpk
        INNER JOIN stockgroup g ON g.cGRPpk = s.cSTKfkGRP
        WHERE s.nSTKsuspend = 0
          AND sd.cSTDcode IS NOT NULL
          AND TRIM(sd.cSTDcode) <> ''
        GROUP BY s.cSTKpk, TRIM(s.cSTKdesc), TRIM(g.cGRPdesc)
    )
"""

RINGKASAN_FIELDS = [
    ("golongan", "Golongan"),
    ("kode_barang", "Kode"),
    ("nama_barang", "Nama Barang"),
    ("total_masuk", "Total Masuk"),
    ("total_keluar", "Total Keluar"),
    ("net_stok", "Net Stok"),
    ("jumlah_transaksi", "Jumlah Transaksi"),
]

TRX_FIELDS = [
    ("kode_barang", "Kode"),
    ("nama_barang", "Nama Barang"),
    ("golongan", "Golongan"),
    ("tanggal", "Tanggal"),
    ("no_faktur", "No Faktur"),
    ("jenis_transaksi", "Jenis Transaksi"),
    ("pihak", "Supplier/Pelanggan"),
    ("qty_masuk", "Qty Masuk"),
    ("qty_keluar", "Qty Keluar"),
    ("mutasi", "Mutasi"),
    ("harga", "Harga"),
    ("cSTKpk", "cSTKpk"),
    ("cIVDpk", "cIVDpk"),
]

GABUNGAN_FIELDS = [
    ("kode_barang", "Kode"),
    ("nama_barang", "Nama Barang"),
    ("golongan", "Golongan"),
    ("net_stok_ringkasan", "Net Stok (Ringkasan)"),
    ("urutan", "Urutan"),
    ("total_transaksi", "Total Transaksi"),
    ("tanggal", "Tanggal"),
    ("no_faktur", "No Faktur"),
    ("jenis_transaksi", "Jenis Transaksi"),
    ("pihak", "Supplier/Pelanggan"),
    ("qty_masuk", "Qty Masuk"),
    ("qty_keluar", "Qty Keluar"),
    ("mutasi", "Mutasi"),
    ("saldo_berjalan", "Saldo Berjalan"),
    ("cocok_akhir", "Cocok Akhir"),
]


def _table_columns(mysql, table: str) -> dict[str, str]:
    with mysql.cursor() as cur:
        cur.execute(f"SHOW COLUMNS FROM `{table}`")
        rows = cur.fetchall()
    return {str(r["Field"]).lower(): str(r["Field"]) for r in rows}


def _pick(cols: dict[str, str], *candidates: str) -> str | None:
    for name in candidates:
        hit = cols.get(name.lower())
        if hit:
            return hit
    return None


def _invoice_select_parts(mysql) -> tuple[str, str, str, str, str]:
    """Return (no_faktur_expr, jenis_join_sql, jenis_expr, pihak_join_sql, pihak_expr)."""
    inv = _table_columns(mysql, "invoice")
    no_col = _pick(inv, "cINVrefno", "cINVno", "cINVref", "cINVpk")
    no_expr = f"TRIM(COALESCE(i.`{no_col}`, ''))" if no_col else "''"

    frm_col = _pick(inv, "cINVfkEXC", "cINVfkFRM", "cINVfkFMT", "cINVtype")
    ent_col = _pick(inv, "cINVfkENT")

    jenis_join = ""
    jenis_expr = "''"
    if frm_col:
        for table in ("formula", "formulatype", "formulamaster"):
            try:
                frm = _table_columns(mysql, table)
            except Exception:
                continue
            pk = _pick(frm, "cINVpk", "cFRMpk", "cFMTpk")
            desc = _pick(frm, "cINVremark", "serino", "cFRMdesc", "cFMTdesc", "cFRMname")
            if pk and desc:
                jenis_join = f"LEFT JOIN `{table}` f ON f.`{pk}` = i.`{frm_col}`"
                jenis_expr = f"TRIM(COALESCE(f.`{desc}`, ''))"
                break

    pihak_join = ""
    pihak_expr = "''"
    if ent_col:
        try:
            ent = _table_columns(mysql, "entity")
            ent_pk = _pick(ent, "cENTpk")
            ent_desc = _pick(ent, "cENTdesc", "cENTname")
            if ent_pk and ent_desc:
                pihak_join = f"LEFT JOIN entity e ON e.`{ent_pk}` = i.`{ent_col}`"
                pihak_expr = f"TRIM(COALESCE(e.`{ent_desc}`, ''))"
        except Exception:
            pass

    return no_expr, jenis_join, jenis_expr, pihak_join, pihak_expr


def ringkasan_sql(golongan: str) -> str:
    return f"""
        WITH {BARANG_KODE_CTE}
        SELECT
            bk.golongan,
            bk.kode_barang,
            bk.nama_barang,
            SUM(COALESCE(d.nIVDqtyin, 0)) AS total_masuk,
            SUM(COALESCE(d.nIVDqtyout, 0)) AS total_keluar,
            SUM(COALESCE(d.nIVDqtyin, 0) - COALESCE(d.nIVDqtyout, 0)) AS net_stok,
            COUNT(*) AS jumlah_transaksi
        FROM invoicedetail d
        INNER JOIN barang_kode bk ON bk.cSTKpk = d.cIVDfkSTK
        WHERE bk.golongan = %s
        GROUP BY bk.golongan, bk.kode_barang, bk.nama_barang
        HAVING net_stok < 0
        ORDER BY bk.kode_barang
    """


def transaksi_sql(
    *,
    no_faktur_expr: str,
    jenis_join: str,
    jenis_expr: str,
    pihak_join: str,
    pihak_expr: str,
) -> str:
    return f"""
        WITH {BARANG_KODE_CTE},
        ringkasan AS (
            SELECT
                bk.golongan,
                bk.kode_barang,
                bk.nama_barang,
                SUM(COALESCE(d.nIVDqtyin, 0) - COALESCE(d.nIVDqtyout, 0)) AS net_stok
            FROM invoicedetail d
            INNER JOIN barang_kode bk ON bk.cSTKpk = d.cIVDfkSTK
            WHERE bk.golongan = %s
            GROUP BY bk.golongan, bk.kode_barang, bk.nama_barang
            HAVING net_stok < 0
        )
        SELECT
            bk.kode_barang,
            bk.nama_barang,
            bk.golongan,
            i.dINVdate AS tanggal,
            {no_faktur_expr} AS no_faktur,
            {jenis_expr} AS jenis_transaksi,
            {pihak_expr} AS pihak,
            COALESCE(d.nIVDqtyin, 0) AS qty_masuk,
            COALESCE(d.nIVDqtyout, 0) AS qty_keluar,
            COALESCE(d.nIVDqtyin, 0) - COALESCE(d.nIVDqtyout, 0) AS mutasi,
            COALESCE(d.nIVDprice, 0) AS harga,
            bk.cSTKpk,
            d.cIVDpk
        FROM invoicedetail d
        INNER JOIN invoice i ON i.cINVpk = d.cIVDfkINV
        INNER JOIN barang_kode bk ON bk.cSTKpk = d.cIVDfkSTK
        INNER JOIN ringkasan r ON r.kode_barang = bk.kode_barang
        {jenis_join}
        {pihak_join}
        ORDER BY bk.kode_barang, i.dINVdate, no_faktur, d.cIVDpk
    """


def _fmt_date(value) -> str:
    if value is None or value == "":
        return ""
    if isinstance(value, datetime):
        return value.strftime("%Y-%m-%d %H:%M")
    if isinstance(value, date):
        return value.isoformat()
    return str(value)


def _build_gabungan(
    ringkasan: list[dict],
    transaksi: list[dict],
) -> list[dict]:
    net_by_kode = {str(r["kode_barang"]): int(r["net_stok"]) for r in ringkasan}
    by_kode: dict[str, list[dict]] = {}
    for row in transaksi:
        kode = str(row["kode_barang"])
        by_kode.setdefault(kode, []).append(row)

    gabungan: list[dict] = []
    for kode in sorted(by_kode.keys()):
        rows = by_kode[kode]
        net = net_by_kode.get(kode, 0)
        saldo = 0
        total = len(rows)
        for idx, row in enumerate(rows, start=1):
            mutasi = int(row.get("mutasi") or 0)
            saldo += mutasi
            cocok = "YA" if idx == total and saldo == net else ""
            gabungan.append(
                {
                    "kode_barang": row.get("kode_barang"),
                    "nama_barang": row.get("nama_barang"),
                    "golongan": row.get("golongan"),
                    "net_stok_ringkasan": net,
                    "urutan": idx,
                    "total_transaksi": total,
                    "tanggal": _fmt_date(row.get("tanggal")),
                    "no_faktur": row.get("no_faktur"),
                    "jenis_transaksi": row.get("jenis_transaksi"),
                    "pihak": row.get("pihak"),
                    "qty_masuk": row.get("qty_masuk"),
                    "qty_keluar": row.get("qty_keluar"),
                    "mutasi": mutasi,
                    "saldo_berjalan": saldo,
                    "cocok_akhir": cocok,
                }
            )
    return gabungan


def _num(value) -> float | int | None:
    if value is None or value == "":
        return None
    try:
        n = float(value)
        return int(n) if n == int(n) else n
    except (TypeError, ValueError):
        return None


def _style_sheet(ws, header_count: int, *, total_row: int | None = None) -> None:
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

    header_fill = PatternFill("solid", fgColor="2F5496")
    header_font = Font(bold=True, color="FFFFFF", size=11)
    total_fill = PatternFill("solid", fgColor="E2EFDA")
    total_font = Font(bold=True, size=11)
    thin = Side(style="thin", color="D9D9D9")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)

    ws.row_dimensions[1].height = 30
    for col in range(1, header_count + 1):
        cell = ws.cell(row=1, column=col)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = border

    max_data_row = total_row or ws.max_row
    for row in ws.iter_rows(min_row=2, max_row=max_data_row, min_col=1, max_col=header_count):
        for cell in row:
            cell.border = border
            cell.alignment = Alignment(vertical="center", wrap_text=False)
            if isinstance(cell.value, (int, float)):
                cell.alignment = Alignment(horizontal="right", vertical="center")
                if isinstance(cell.value, float) and cell.value == int(cell.value):
                    cell.value = int(cell.value)

    if total_row:
        for col in range(1, header_count + 1):
            cell = ws.cell(row=total_row, column=col)
            cell.fill = total_fill
            cell.font = total_font
            cell.border = border

    for col_cells in ws.columns:
        letter = col_cells[0].column_letter
        width = 12
        for cell in col_cells:
            if cell.value is not None:
                text = str(cell.value)
                width = max(width, min(len(text) + 3, 50))
        ws.column_dimensions[letter].width = width

    ws.freeze_panes = "A2"
    data_last = (total_row - 1) if total_row else ws.max_row
    if data_last > 1:
        ws.auto_filter.ref = f"A1:{ws.cell(1, header_count).column_letter}{data_last}"


def _write_data_sheet(
    ws,
    fields: list[tuple[str, str]],
    rows: list[dict],
    *,
    sum_keys: list[str] | None = None,
    total_label_key: str | None = None,
    total_label: str = "TOTAL",
    extra_total: dict[str, str] | None = None,
) -> None:
    keys = [k for k, _ in fields]
    labels = [label for _, label in fields]
    ws.append(labels)

    totals: dict[str, float] = {k: 0.0 for k in (sum_keys or [])}
    for row in rows:
        line: list = []
        for key in keys:
            value = row.get(key)
            if key == "tanggal":
                value = _fmt_date(value)
            line.append(value)
            if key in totals:
                n = _num(value)
                if n is not None:
                    totals[key] += n
        ws.append(line)

    if sum_keys:
        total_row_idx = ws.max_row + 1
        label_col = keys.index(total_label_key) + 1 if total_label_key and total_label_key in keys else 1
        total_line: list = [""] * len(keys)
        total_line[label_col - 1] = total_label
        for key in sum_keys:
            if key in keys:
                val = totals[key]
                total_line[keys.index(key)] = int(val) if val == int(val) else val
        if extra_total:
            for key, text in extra_total.items():
                if key in keys:
                    total_line[keys.index(key)] = text
        ws.append(total_line)
        _style_sheet(ws, len(labels), total_row=total_row_idx)
    else:
        _style_sheet(ws, len(labels))


def _write_info_sheet(ws, *, db_name: str, golongan: str, stats: dict) -> None:
    from openpyxl.styles import Font

    ws.column_dimensions["A"].width = 32
    ws.column_dimensions["B"].width = 60
    ws["A1"] = "Export Bukti Stok Minus (MySQL Desktop)"
    ws["A1"].font = Font(bold=True, size=14)
    lines = [
        ("Database", db_name),
        ("Golongan", golongan),
        ("Barang minus", stats.get("barang_minus", "")),
        ("Baris transaksi", stats.get("transaksi_rows", "")),
        ("Baris gabungan (ledger)", stats.get("gabungan_rows", "")),
        ("Cocok saldo akhir", f"{stats.get('cocok_akhir', 0)}/{stats.get('barang_minus', 0)}"),
        ("", ""),
        ("Sheet Ringkasan Minus", "1 baris per kode barang — net stok < 0"),
        ("Sheet Transaksi", "Detail faktur per barang minus"),
        ("Sheet Gabungan Bukti", "Ledger + saldo berjalan; Cocok Akhir = YA"),
    ]
    for idx, (label, value) in enumerate(lines, start=2):
        ws.cell(row=idx, column=1, value=label)
        ws.cell(row=idx, column=2, value=value)


def export_stok_minus_bukti_xlsx(
    mysql,
    *,
    golongan: str = DEFAULT_GOLONGAN,
    out_path: Path | None = None,
    db_name: str = "",
) -> dict:
    out = (out_path or default_output_path(golongan, db_name)).expanduser().resolve()
    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)

    no_expr, jenis_join, jenis_expr, pihak_join, pihak_expr = _invoice_select_parts(mysql)

    with Spinner(f"Ringkasan stok minus — {golongan}..."):
        ringkasan = fetch_all_dict(mysql, ringkasan_sql(golongan), (golongan,))

    with Spinner(f"Detail transaksi ({len(ringkasan)} barang minus)..."):
        trx_sql = transaksi_sql(
            no_faktur_expr=no_expr,
            jenis_join=jenis_join,
            jenis_expr=jenis_expr,
            pihak_join=pihak_join,
            pihak_expr=pihak_expr,
        )
        transaksi = fetch_all_dict(mysql, trx_sql, (golongan,))

    gabungan = _build_gabungan(ringkasan, transaksi)
    cocok = sum(1 for r in gabungan if r.get("cocok_akhir") == "YA")
    mismatch = len(ringkasan) - cocok

    stats = {
        "barang_minus": len(ringkasan),
        "transaksi_rows": len(transaksi),
        "gabungan_rows": len(gabungan),
        "cocok_akhir": cocok,
    }

    from openpyxl import Workbook

    wb = Workbook()
    info = wb.active
    info.title = "Info"
    _write_info_sheet(info, db_name=db_name, golongan=golongan, stats=stats)

    ws_ring = wb.create_sheet("Ringkasan Minus")
    _write_data_sheet(
        ws_ring,
        RINGKASAN_FIELDS,
        ringkasan,
        sum_keys=["total_masuk", "total_keluar", "net_stok", "jumlah_transaksi"],
        total_label_key="kode_barang",
        total_label=f"TOTAL ({len(ringkasan)} barang)",
    )

    ws_trx = wb.create_sheet("Transaksi")
    _write_data_sheet(
        ws_trx,
        TRX_FIELDS,
        transaksi,
        sum_keys=["qty_masuk", "qty_keluar", "mutasi"],
        total_label_key="kode_barang",
        total_label=f"TOTAL ({len(transaksi)} baris)",
    )

    ws_gab = wb.create_sheet("Gabungan Bukti")
    _write_data_sheet(
        ws_gab,
        GABUNGAN_FIELDS,
        gabungan,
        sum_keys=["qty_masuk", "qty_keluar", "mutasi"],
        total_label_key="kode_barang",
        total_label=f"TOTAL ({len(gabungan)} baris)",
        extra_total={"cocok_akhir": f"{cocok} YA"},
    )

    wb.save(out)

    return {
        "golongan": golongan,
        "db_name": db_name,
        "barang_minus": len(ringkasan),
        "transaksi_rows": len(transaksi),
        "gabungan_rows": len(gabungan),
        "cocok_akhir": cocok,
        "mismatch_akhir": mismatch,
        "out_path": str(out),
    }
