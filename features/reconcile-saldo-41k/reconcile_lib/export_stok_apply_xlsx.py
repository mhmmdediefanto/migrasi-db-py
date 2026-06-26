from __future__ import annotations

import csv
from datetime import datetime
from pathlib import Path

from openpyxl import Workbook

from reconcile_lib.export_xlsx import (
    BORDER,
    HEADER_FILL,
    LEFT,
    RIGHT,
    SUBTITLE_FONT,
    TITLE_FONT,
    TOTAL_FILL,
    TOTAL_FONT,
    _apply_sheet_layout,
    _autosize_columns,
    _style_data_area,
    _style_header_row,
)
from reconcile_lib.stok_apply import PREVIEW_FIELDS

APPLY_HEADERS = [
    ("skenario", "Skenario"),
    ("kode_barang", "Kode Barang"),
    ("nama_barang", "Nama Barang"),
    ("stok_sebelum", "Stok Sebelum"),
    ("stok_target", "Stok Target"),
    ("delta", "Delta"),
    ("aksi", "Aksi"),
    ("catatan", "Catatan"),
]

NUMERIC_KEYS = {"stok_sebelum", "stok_target", "delta"}


def _load_preview_csv(path: Path) -> list[dict]:
    with path.open(encoding="utf-8", newline="") as f:
        return list(csv.DictReader(f))


def _write_apply_sheet(wb: Workbook, title: str, rows: list[dict]) -> None:
    keys = [k for k, _ in APPLY_HEADERS]
    headers = [label for _, label in APPLY_HEADERS]
    ws = wb.create_sheet(title)
    ws.append(headers)
    _style_header_row(ws, 1, len(headers))
    num_cols = {i + 1 for i, k in enumerate(keys) if k in NUMERIC_KEYS}

    for row in rows:
        ws.append([
            int(row[k]) if k in NUMERIC_KEYS and row.get(k) not in ("", None) else row.get(k, "")
            for k in keys
        ])

    if rows:
        _style_data_area(ws, 2, ws.max_row, len(headers), num_cols)
        tr = ws.max_row + 1
        ws.cell(row=tr, column=1, value="TOTAL")
        ws.cell(row=tr, column=4, value=sum(int(r.get("stok_sebelum") or 0) for r in rows))
        ws.cell(row=tr, column=5, value=sum(int(r.get("stok_target") or 0) for r in rows))
        ws.cell(row=tr, column=6, value=sum(int(r.get("delta") or 0) for r in rows))
        for col in range(1, len(headers) + 1):
            c = ws.cell(row=tr, column=col)
            c.fill = TOTAL_FILL
            c.font = TOTAL_FONT
            c.border = BORDER
            c.alignment = RIGHT if col in num_cols else LEFT

    ws.freeze_panes = "A2"
    if rows:
        ws.auto_filter.ref = f"A1:H{ws.max_row - 1}"
    _autosize_columns(ws, len(headers), start_row=1)
    _apply_sheet_layout(ws, len(headers))


def export_stok_apply_workbook(
    path: Path,
    rows: list[dict],
    stats: dict,
) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    wb = Workbook()
    ws = wb.active
    ws.title = "Ringkasan"
    ws["A1"] = "Koreksi Stok S1/S2/S3 — matahari_final"
    ws["A1"].font = TITLE_FONT
    ws.merge_cells("A1:D1")

    meta = [
        ("S1", "Luar 41k, tanpa transaksi → stok = 0"),
        ("S2", "Luar 41k, ada transaksi → stok = 0 + net trx"),
        ("S3", "Dalam 41k → stok = Excel 18/6 + net trx sejak 18/6"),
        ("Database", stats.get("pg_database", "")),
        ("Gudang", stats.get("gudang", "")),
        ("Sumber Excel", stats.get("source_file", "")),
        ("Dibuat", datetime.now().strftime("%Y-%m-%d %H:%M")),
    ]
    row = 3
    for label, value in meta:
        ws.cell(row=row, column=1, value=label).font = SUBTITLE_FONT
        ws.cell(row=row, column=2, value=value)
        row += 1

    row += 1
    ws.cell(row=row, column=1, value="Metrik").font = SUBTITLE_FONT
    ws.cell(row=row, column=2, value="Jumlah").font = SUBTITLE_FONT
    _style_header_row(ws, row, 2)
    row += 1
    start = row
    summary_rows = [
        ("s1_kandidat", "S1 kandidat (luar 41k tanpa trx)"),
        ("s1_apply", "S1 di-apply"),
        ("s2_apply", "S2 di-apply"),
        ("s3_apply", "S3 di-apply"),
        ("total_apply", "Total baris koreksi"),
        ("update", "Update stok_akhir"),
        ("insert", "Insert stok_akhir baru"),
    ]
    for key, label in summary_rows:
        ws.cell(row=row, column=1, value=label)
        ws.cell(row=row, column=2, value=stats.get(key, 0))
        row += 1
    _style_data_area(ws, start, row - 1, 2, {2})
    ws.column_dimensions["A"].width = 44
    ws.column_dimensions["B"].width = 18

    by_skenario = {
        "S1": [r for r in rows if r.get("skenario") == "S1"],
        "S2": [r for r in rows if r.get("skenario") == "S2"],
        "S3": [r for r in rows if r.get("skenario") == "S3"],
    }
    _write_apply_sheet(wb, "Koreksi S1", by_skenario["S1"])
    _write_apply_sheet(wb, "Koreksi S2", by_skenario["S2"])
    _write_apply_sheet(wb, "Koreksi S3", by_skenario["S3"])
    _write_apply_sheet(wb, "Semua Koreksi", rows)

    wb.save(path)


def export_stok_apply_from_csv(
    csv_path: Path,
    xlsx_path: Path,
    stats: dict | None = None,
) -> dict:
    rows = _load_preview_csv(csv_path)
    inner = stats or {}
    inner.setdefault("pg_database", "matahari_final")
    inner.setdefault("s1_apply", sum(1 for r in rows if r.get("skenario") == "S1"))
    inner.setdefault("s2_apply", sum(1 for r in rows if r.get("skenario") == "S2"))
    inner.setdefault("s3_apply", sum(1 for r in rows if r.get("skenario") == "S3"))
    inner.setdefault("total_apply", len(rows))
    inner.setdefault("update", sum(1 for r in rows if r.get("aksi") == "update"))
    inner.setdefault("insert", sum(1 for r in rows if r.get("aksi") == "insert"))
    inner.setdefault("s1_kandidat", inner["s1_apply"])
    export_stok_apply_workbook(xlsx_path, rows, inner)
    return {"rows": len(rows), **inner}
