from __future__ import annotations

from datetime import datetime
from pathlib import Path

from openpyxl import Workbook

from reconcile_lib.export_xlsx import (
    BORDER,
    HEADER_FILL,
    HEADER_FONT,
    LEFT,
    RIGHT,
    SUBTITLE_FONT,
    TITLE_FONT,
    TOTAL_FILL,
    TOTAL_FONT,
    _apply_sheet_layout,
    _autosize_columns,
    _style_data_area,
    _style_header_row,
)
from reconcile_lib.stok_hitung_query import OUTPUT_FIELDS

SUMMARY_META = [
    ("total_41k", "Baris Excel 41k"),
    ("total_luar_41k", "Barang luar 41k (stok/trx)"),
    ("selisih_gudang_count", "Total selisih vs stok gudang"),
    ("selisih_gudang_41k", "Selisih gudang — dalam 41k"),
    ("selisih_gudang_luar", "Selisih gudang — luar 41k"),
    ("ada_trx_41k", "Ada transaksi (41k)"),
    ("ada_trx_luar", "Ada transaksi (luar 41k)"),
    ("total_baseline_41k", "Total baseline Excel"),
    ("total_stok_hitung_41k", "Total stok hitung (41k)"),
    ("total_stok_hitung_luar", "Total stok hitung (luar 41k)"),
]

NUMERIC_KEYS = {
    "baseline", "net_pembelian", "net_penjualan", "net_retur_beli", "net_retur_jual",
    "net_transfer_masuk", "net_transfer_keluar", "net_penyesuaian",
    "stok_hitung", "stok_gudang_sistem", "stok_barang_sistem", "selisih_gudang", "selisih_barang",
}


def _write_detail_sheet(wb: Workbook, title: str, rows: list[dict], *, only_selisih: bool = False) -> None:
    keys = [k for k, _ in OUTPUT_FIELDS]
    headers = [label for _, label in OUTPUT_FIELDS]
    data = rows
    if only_selisih:
        data = [r for r in rows if r.get("stok_gudang_sistem") != "" and int(r.get("selisih_gudang") or 0) != 0]

    ws = wb.create_sheet(title)
    ws.append(headers)
    _style_header_row(ws, 1, len(headers))
    num_cols = {i + 1 for i, k in enumerate(keys) if k in NUMERIC_KEYS or k == "no"}

    for src in data:
        ws.append([src.get(k, "") for k in keys])

    if data:
        _style_data_area(ws, 2, ws.max_row, len(headers), num_cols)
        tr = ws.max_row + 1
        ws.cell(row=tr, column=1, value="TOTAL")
        ws.cell(row=tr, column=5, value=sum(int(r.get("baseline") or 0) for r in data))
        ws.cell(row=tr, column=13, value=sum(int(r.get("stok_hitung") or 0) for r in data))
        ws.cell(row=tr, column=14, value=sum(int(r.get("stok_gudang_sistem") or 0) for r in data if r.get("stok_gudang_sistem") != ""))
        ws.cell(row=tr, column=16, value=sum(int(r.get("selisih_gudang") or 0) for r in data if r.get("stok_gudang_sistem") != ""))
        for col in range(1, len(headers) + 1):
            c = ws.cell(row=tr, column=col)
            c.fill = TOTAL_FILL
            c.font = TOTAL_FONT
            c.border = BORDER
            c.alignment = RIGHT if col in num_cols else LEFT

    ws.freeze_panes = "A2"
    if data:
        ws.auto_filter.ref = f"A1:{chr(64+len(headers)) if len(headers)<=26 else 'Q'}{ws.max_row - (1 if data else 0)}"
    _autosize_columns(ws, len(headers), start_row=1)
    _apply_sheet_layout(ws, len(headers))


def export_stok_hitung_workbook(
    path: Path,
    rows_41k: list[dict],
    rows_luar: list[dict],
    stats: dict,
) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    wb = Workbook()
    ws = wb.active
    ws.title = "Ringkasan"
    ws["A1"] = "Stok Hitung — Baseline Excel/0 ± Transaksi (abaikan stok sistem)"
    ws["A1"].font = TITLE_FONT
    ws.merge_cells("A1:D1")

    meta = [
        ("Rumus", "stok_hitung = baseline (Excel jika dalam 41k, else 0) ± transaksi web"),
        ("Trx 41k sejak", stats.get("since_date", "")),
        ("Trx luar 41k", "SEMUA transaksi web (baseline=0)"),
        ("Database", stats.get("pg_database", "")),
        ("Gudang", stats.get("gudang", "")),
        ("Sumber Excel", stats.get("source_file", "")),
        ("Dibuat", datetime.now().strftime("%Y-%m-%d %H:%M")),
    ]
    row = 3
    for label, value in meta:
        ws.cell(row=row, column=1, value=label).font = SUBTITLE_FONT
        ws.cell(row=row, column=2, value=value)
        row += 1

    row += 1
    ws.cell(row=row, column=1, value="Metrik").font = SUBTITLE_FONT
    ws.cell(row=row, column=2, value="Jumlah").font = SUBTITLE_FONT
    _style_header_row(ws, row, 2)
    row += 1
    start = row
    for key, label in SUMMARY_META:
        ws.cell(row=row, column=1, value=label)
        ws.cell(row=row, column=2, value=stats.get(key, 0))
        row += 1
    _style_data_area(ws, start, row - 1, 2, {2})
    ws.column_dimensions["A"].width = 44
    ws.column_dimensions["B"].width = 18

    _write_detail_sheet(wb, "Hitung 41k", rows_41k)
    _write_detail_sheet(wb, "Selisih 41k", rows_41k, only_selisih=True)
    _write_detail_sheet(wb, "Hitung Luar 41k", rows_luar)
    _write_detail_sheet(wb, "Selisih Luar 41k", rows_luar, only_selisih=True)

    wb.save(path)
