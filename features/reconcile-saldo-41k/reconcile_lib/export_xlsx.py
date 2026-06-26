from __future__ import annotations

from datetime import datetime
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from reconcile_lib.trx_query import OUTPUT_FIELDS, SUMMARY_FIELDS

HEADER_FILL = PatternFill("solid", fgColor="1F4E79")
HEADER_FONT = Font(bold=True, color="FFFFFF", size=11)
TITLE_FONT = Font(bold=True, size=14, color="1F4E79")
SUBTITLE_FONT = Font(bold=True, size=11, color="333333")
TOTAL_FILL = PatternFill("solid", fgColor="D9E1F2")
TOTAL_FONT = Font(bold=True, size=11)
THIN = Side(style="thin", color="BFBFBF")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
CENTER = Alignment(horizontal="center", vertical="center", wrap_text=True)
LEFT = Alignment(horizontal="left", vertical="center", wrap_text=True)
RIGHT = Alignment(horizontal="right", vertical="center")


def _cell_width(value) -> float:
    if value is None:
        return 10
    text = str(value)
    length = max(len(text), 8)
    return min(max(length * 1.15 + 2, 10), 55)


def _format_value(key: str, value):
    if key.startswith("ada_") or key in ("not_in_barang", "ada_trx_apapun"):
        if isinstance(value, bool):
            return "Y" if value else "N"
        return value
    if key.startswith("cnt_") or key in ("stok_excel", "stok_barang_system", "stok_gudang_pusat", "no"):
        if value is None or value == "":
            return 0 if key != "no" else ""
        try:
            return int(round(float(value)))
        except (TypeError, ValueError):
            return value
    return value if value is not None else ""


def _apply_sheet_layout(sheet, header_cols: int, data_start: int = 2) -> None:
    sheet.sheet_format.defaultRowHeight = 18
    sheet.row_dimensions[1].height = 22
    for row in range(data_start, sheet.max_row + 1):
        sheet.row_dimensions[row].height = 18
    for col in range(1, header_cols + 1):
        letter = get_column_letter(col)
        current = sheet.column_dimensions[letter].width or 10
        sheet.column_dimensions[letter].width = max(current, 12)


def _autosize_columns(sheet, col_count: int, start_row: int = 1) -> None:
    widths: dict[int, float] = {}
    for row in sheet.iter_rows(min_row=start_row, max_row=sheet.max_row, max_col=col_count):
        for idx, cell in enumerate(row, start=1):
            widths[idx] = max(widths.get(idx, 10), _cell_width(cell.value))
    for idx, width in widths.items():
        letter = get_column_letter(idx)
        sheet.column_dimensions[letter].width = width


def _style_header_row(sheet, row_num: int, col_count: int) -> None:
    for col in range(1, col_count + 1):
        cell = sheet.cell(row=row_num, column=col)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = CENTER
        cell.border = BORDER


def _style_data_area(sheet, start_row: int, end_row: int, col_count: int, numeric_cols: set[int]) -> None:
    for row in range(start_row, end_row + 1):
        for col in range(1, col_count + 1):
            cell = sheet.cell(row=row, column=col)
            cell.border = BORDER
            cell.alignment = RIGHT if col in numeric_cols else LEFT


def export_workbook(path: Path, rows: list[dict], stats: dict) -> dict:
    path.parent.mkdir(parents=True, exist_ok=True)
    wb = Workbook()

    # --- Sheet Ringkasan ---
    ws_sum = wb.active
    ws_sum.title = "Ringkasan"
    ws_sum["A1"] = "Cek Transaksi — Saldo Stock 41k vs matahari_final"
    ws_sum["A1"].font = TITLE_FONT
    ws_sum.merge_cells("A1:D1")

    meta = [
        ("Sumber Excel", stats.get("source_file", "")),
        ("Database", stats.get("pg_database", "")),
        ("Gudang", stats.get("gudang", "")),
        ("Filter Tanggal Trx", stats.get("since_date", "SEMUA")),
        ("Dibuat", datetime.now().strftime("%Y-%m-%d %H:%M")),
    ]
    row = 3
    for label, value in meta:
        ws_sum.cell(row=row, column=1, value=label).font = SUBTITLE_FONT
        ws_sum.cell(row=row, column=2, value=value)
        row += 1

    row += 1
    ws_sum.cell(row=row, column=1, value="Metrik").font = SUBTITLE_FONT
    ws_sum.cell(row=row, column=2, value="Jumlah").font = SUBTITLE_FONT
    _style_header_row(ws_sum, row, 2)
    row += 1
    summary_start = row
    for key, label in SUMMARY_FIELDS:
        ws_sum.cell(row=row, column=1, value=label)
        ws_sum.cell(row=row, column=2, value=stats.get(key, 0))
        row += 1
    _style_data_area(ws_sum, summary_start, row - 1, 2, {2})
    _autosize_columns(ws_sum, 2, start_row=1)
    ws_sum.column_dimensions["A"].width = max(ws_sum.column_dimensions["A"].width, 38)
    ws_sum.column_dimensions["B"].width = max(ws_sum.column_dimensions["B"].width, 18)

    # --- Sheet Detail ---
    ws = wb.create_sheet("Detail Trx 41k")
    headers = [label for _, label in OUTPUT_FIELDS]
    keys = [key for key, _ in OUTPUT_FIELDS]

    ws.append(headers)
    _style_header_row(ws, 1, len(headers))

    numeric_keys = {
        "no",
        "stok_excel",
        "stok_barang_system",
        "stok_gudang_pusat",
        "cnt_pembelian",
        "cnt_penjualan",
        "cnt_transfer",
        "cnt_retur_beli",
        "cnt_retur_jual",
        "cnt_penyesuaian",
    }
    numeric_cols = {i + 1 for i, key in enumerate(keys) if key in numeric_keys}

    for src in rows:
        ws.append([_format_value(key, src.get(key)) for key in keys])

    data_end = ws.max_row
    _style_data_area(ws, 2, data_end, len(headers), numeric_cols)

    # Baris TOTAL
    total_row = data_end + 1
    ws.cell(row=total_row, column=1, value="TOTAL")
    ws.cell(row=total_row, column=4, value=stats.get("total_stok_excel", 0))
    ws.cell(row=total_row, column=5, value=stats.get("total_stok_system", 0))
    ws.cell(row=total_row, column=6, value=stats.get("total_stok_gudang", 0))

    count_map = {
        7: stats.get("ada_pembelian", 0),
        8: sum(int(r.get("cnt_pembelian") or 0) for r in rows),
        10: stats.get("ada_penjualan", 0),
        11: sum(int(r.get("cnt_penjualan") or 0) for r in rows),
        14: stats.get("ada_transfer", 0),
        15: sum(int(r.get("cnt_transfer") or 0) for r in rows),
        16: stats.get("ada_retur_beli", 0),
        17: sum(int(r.get("cnt_retur_beli") or 0) for r in rows),
        18: stats.get("ada_retur_jual", 0),
        19: sum(int(r.get("cnt_retur_jual") or 0) for r in rows),
        20: stats.get("ada_penyesuaian", 0),
        21: sum(int(r.get("cnt_penyesuaian") or 0) for r in rows),
        22: stats.get("ada_trx", 0),
        24: stats.get("not_in_barang", 0),
    }
    for col, val in count_map.items():
        ws.cell(row=total_row, column=col, value=val)

    for col in range(1, len(headers) + 1):
        cell = ws.cell(row=total_row, column=col)
        cell.fill = TOTAL_FILL
        cell.font = TOTAL_FONT
        cell.border = BORDER
        cell.alignment = RIGHT if col in numeric_cols or col >= 4 else LEFT

    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(headers))}{data_end}"
    _autosize_columns(ws, len(headers), start_row=1)
    _apply_sheet_layout(ws, len(headers))

    # --- Sheet Tanpa Trx ---
    ws0 = wb.create_sheet("Tanpa Trx")
    no_trx = [r for r in rows if not r.get("ada_trx_apapun") and not r.get("not_in_barang")]
    ws0.append(headers)
    _style_header_row(ws0, 1, len(headers))
    for src in no_trx:
        ws0.append([_format_value(key, src.get(key)) for key in keys])
    if no_trx:
        end0 = ws0.max_row
        _style_data_area(ws0, 2, end0, len(headers), numeric_cols)
        tr = end0 + 1
        ws0.cell(row=tr, column=1, value="TOTAL")
        ws0.cell(row=tr, column=2, value=f"{len(no_trx)} barang")
        ws0.cell(row=tr, column=4, value=sum(int(r.get("stok_excel") or 0) for r in no_trx))
        ws0.cell(row=tr, column=5, value=sum(int(r.get("stok_barang_system") or 0) for r in no_trx if r.get("stok_barang_system") is not None))
        ws0.cell(row=tr, column=6, value=sum(int(r.get("stok_gudang_pusat") or 0) for r in no_trx if r.get("stok_gudang_pusat") is not None))
        for col in range(1, len(headers) + 1):
            c = ws0.cell(row=tr, column=col)
            c.fill = TOTAL_FILL
            c.font = TOTAL_FONT
            c.border = BORDER
    ws0.freeze_panes = "A2"
    _autosize_columns(ws0, len(headers), start_row=1)
    _apply_sheet_layout(ws0, len(headers))

    # --- Sheet Ada Trx ---
    ws1 = wb.create_sheet("Ada Trx")
    ada_trx = [r for r in rows if r.get("ada_trx_apapun")]
    ws1.append(headers)
    _style_header_row(ws1, 1, len(headers))
    for src in ada_trx:
        ws1.append([_format_value(key, src.get(key)) for key in keys])
    if ada_trx:
        end1 = ws1.max_row
        _style_data_area(ws1, 2, end1, len(headers), numeric_cols)
        tr = end1 + 1
        ws1.cell(row=tr, column=1, value="TOTAL")
        ws1.cell(row=tr, column=2, value=f"{len(ada_trx)} barang")
        ws1.cell(row=tr, column=4, value=sum(int(r.get("stok_excel") or 0) for r in ada_trx))
        ws1.cell(row=tr, column=5, value=sum(int(r.get("stok_barang_system") or 0) for r in ada_trx if r.get("stok_barang_system") is not None))
        ws1.cell(row=tr, column=6, value=sum(int(r.get("stok_gudang_pusat") or 0) for r in ada_trx if r.get("stok_gudang_pusat") is not None))
        ws1.cell(row=tr, column=22, value=len(ada_trx))
        for col in range(1, len(headers) + 1):
            c = ws1.cell(row=tr, column=col)
            c.fill = TOTAL_FILL
            c.font = TOTAL_FONT
            c.border = BORDER
    ws1.freeze_panes = "A2"
    _autosize_columns(ws1, len(headers), start_row=1)
    _apply_sheet_layout(ws1, len(headers))

    wb.save(path)
    return stats
