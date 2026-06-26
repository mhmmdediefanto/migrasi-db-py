from __future__ import annotations

from datetime import date, datetime
from pathlib import Path

from openpyxl import Workbook
from openpyxl.utils import get_column_letter

from reconcile_lib.export_xlsx import (
    BORDER,
    CENTER,
    HEADER_FILL,
    HEADER_FONT,
    LEFT,
    RIGHT,
    SUBTITLE_FONT,
    TITLE_FONT,
    TOTAL_FILL,
    TOTAL_FONT,
    _apply_sheet_layout,
    _autosize_columns,
    _cell_width,
    _style_data_area,
    _style_header_row,
)

SUMMARY_META = [
    ("excel_barang_rows", "Baris Excel Saldo Stock 41k"),
    ("barang_luar_dengan_trx", "Barang LUAR 41k dengan transaksi"),
    ("penjualan_faktur_luar", "Faktur penjualan (ada barang luar 41k)"),
    ("penjualan_faktur_murni_luar", "  └ murni luar 41k"),
    ("penjualan_faktur_campuran", "  └ campuran dalam+luar"),
    ("penjualan_line_luar", "Baris detail penjualan luar 41k"),
    ("penjualan_qty_luar", "Total qty penjualan luar 41k"),
    ("pembelian_faktur_luar", "Faktur pembelian pusat (ada barang luar 41k)"),
    ("pembelian_faktur_murni_luar", "  └ murni luar 41k"),
    ("pembelian_faktur_campuran", "  └ campuran dalam+luar"),
    ("pembelian_line_luar", "Baris detail pembelian luar 41k"),
    ("pembelian_qty_luar", "Total qty pembelian luar 41k"),
    ("transfer_line_luar", "Baris transfer gudang luar 41k"),
]

PENJUALAN_FAKTUR_COLS = [
    ("no_faktur", "No Faktur"),
    ("tanggal", "Tanggal"),
    ("kasir", "Kasir"),
    ("toko", "Toko"),
    ("gudang_kasir", "Gudang Kasir"),
    ("kategori_faktur", "Kategori Faktur"),
    ("total_line", "Total Line"),
    ("line_dalam_41k", "Line Dalam 41k"),
    ("line_luar_41k", "Line Luar 41k"),
    ("qty_dalam_41k", "Qty Dalam 41k"),
    ("qty_luar_41k", "Qty Luar 41k"),
]

PENJUALAN_DETAIL_COLS = [
    ("no_faktur", "No Faktur"),
    ("tanggal", "Tanggal"),
    ("kasir", "Kasir"),
    ("toko", "Toko"),
    ("gudang_kasir", "Gudang Kasir"),
    ("kode_barang", "Kode Barang"),
    ("nama_barang", "Nama Barang"),
    ("qty", "Qty"),
    ("harga_satuan", "Harga"),
    ("total", "Total"),
    ("stok_barang_system", "Stok System"),
]

PEMBELIAN_FAKTUR_COLS = [
    ("no_faktur", "No Faktur"),
    ("tanggal", "Tanggal"),
    ("gudang", "Gudang"),
    ("kategori_faktur", "Kategori Faktur"),
    ("total_line", "Total Line"),
    ("line_dalam_41k", "Line Dalam 41k"),
    ("line_luar_41k", "Line Luar 41k"),
    ("qty_dalam_41k", "Qty Dalam 41k"),
    ("qty_luar_41k", "Qty Luar 41k"),
]

PEMBELIAN_DETAIL_COLS = [
    ("no_faktur", "No Faktur"),
    ("tanggal", "Tanggal"),
    ("gudang", "Gudang"),
    ("kode_barang", "Kode Barang"),
    ("nama_barang", "Nama Barang"),
    ("qty", "Qty"),
    ("harga", "Harga"),
    ("total", "Total"),
    ("stok_barang_system", "Stok System"),
]

BARANG_RINGKASAN_COLS = [
    ("kode_barang", "Kode Barang"),
    ("nama_barang", "Nama Barang"),
    ("kategori_trx", "Kategori Trx"),
    ("faktur_jual", "Faktur Jual"),
    ("qty_jual", "Qty Jual"),
    ("tgl_jual_terakhir", "Tgl Jual Terakhir"),
    ("faktur_beli", "Faktur Beli"),
    ("qty_beli", "Qty Beli"),
    ("tgl_beli_terakhir", "Tgl Beli Terakhir"),
    ("stok_barang_system", "Stok System"),
    ("stok_gudang_pusat", "Stok Gudang Pusat"),
]

TRANSFER_COLS = [
    ("transfer_id", "Transfer ID"),
    ("tanggal_transfer", "Tanggal"),
    ("status", "Status"),
    ("gudang_asal", "Gudang Asal"),
    ("gudang_tujuan", "Gudang Tujuan"),
    ("kode_barang", "Kode Barang"),
    ("nama_barang", "Nama Barang"),
    ("jumlah_transfer", "Qty Transfer"),
    ("jumlah_terima", "Qty Terima"),
]


def _fmt_val(value):
    if isinstance(value, datetime):
        return value.strftime("%Y-%m-%d %H:%M")
    if isinstance(value, date):
        return value.isoformat()
    if value is None:
        return ""
    if isinstance(value, float):
        return int(round(value))
    return value


def _numeric_cols(keys: list[str]) -> set[int]:
    numeric = {
        "total_line", "line_dalam_41k", "line_luar_41k", "qty_dalam_41k", "qty_luar_41k",
        "qty", "harga", "harga_satuan", "total", "faktur_jual", "qty_jual", "faktur_beli",
        "qty_beli", "stok_barang_system", "stok_gudang_pusat", "jumlah_transfer", "jumlah_terima",
    }
    return {i + 1 for i, k in enumerate(keys) if k in numeric}


def _write_sheet(
    wb: Workbook,
    title: str,
    columns: list[tuple[str, str]],
    rows: list[dict],
    *,
    total_cols: dict[int, object] | None = None,
) -> None:
    ws = wb.create_sheet(title)
    keys = [k for k, _ in columns]
    headers = [label for _, label in columns]
    num_cols = _numeric_cols(keys)

    ws.append(headers)
    _style_header_row(ws, 1, len(headers))

    for src in rows:
        ws.append([_fmt_val(src.get(k)) for k in keys])

    if rows:
        _style_data_area(ws, 2, ws.max_row, len(headers), num_cols)

    if total_cols:
        tr = ws.max_row + 1
        ws.cell(row=tr, column=1, value="TOTAL")
        for col, val in total_cols.items():
            ws.cell(row=tr, column=col, value=val)
        for col in range(1, len(headers) + 1):
            c = ws.cell(row=tr, column=col)
            c.fill = TOTAL_FILL
            c.font = TOTAL_FONT
            c.border = BORDER
            c.alignment = RIGHT if col in num_cols else LEFT

    ws.freeze_panes = "A2"
    if rows:
        ws.auto_filter.ref = f"A1:{get_column_letter(len(headers))}{ws.max_row - (1 if total_cols else 0)}"
    _autosize_columns(ws, len(headers), start_row=1)
    _apply_sheet_layout(ws, len(headers))


def export_luar_workbook(path: Path, data: dict, stats: dict) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    wb = Workbook()
    ws = wb.active
    ws.title = "Ringkasan"
    ws["A1"] = "Transaksi LUAR Saldo Stock 41k — matahari_final"
    ws["A1"].font = TITLE_FONT
    ws.merge_cells("A1:D1")

    meta = [
        ("Sumber Excel 41k", stats.get("source_file", "")),
        ("Database", stats.get("pg_database", "")),
        ("Gudang pembelian", stats.get("gudang", "")),
        ("Filter tanggal trx", stats.get("since_date", "SEMUA")),
        ("Dibuat", datetime.now().strftime("%Y-%m-%d %H:%M")),
    ]
    row = 3
    for label, value in meta:
        ws.cell(row=row, column=1, value=label).font = SUBTITLE_FONT
        ws.cell(row=row, column=2, value=value)
        row += 1

    row += 1
    ws.cell(row=row, column=1, value="Metrik").font = SUBTITLE_FONT
    ws.cell(row=row, column=2, value="Jumlah").font = SUBTITLE_FONT
    _style_header_row(ws, row, 2)
    row += 1
    start = row
    for key, label in SUMMARY_META:
        ws.cell(row=row, column=1, value=label)
        ws.cell(row=row, column=2, value=stats.get(key, 0))
        row += 1
    _style_data_area(ws, start, row - 1, 2, {2})
    ws.column_dimensions["A"].width = 42
    ws.column_dimensions["B"].width = 18

    pj_f = data["penjualan_faktur"]
    _write_sheet(
        wb, "Penjualan Faktur",
        PENJUALAN_FAKTUR_COLS, pj_f,
        total_cols={
            7: sum(int(r.get("total_line") or 0) for r in pj_f),
            8: sum(int(r.get("line_dalam_41k") or 0) for r in pj_f),
            9: sum(int(r.get("line_luar_41k") or 0) for r in pj_f),
            10: sum(int(r.get("qty_dalam_41k") or 0) for r in pj_f),
            11: sum(int(r.get("qty_luar_41k") or 0) for r in pj_f),
        } if pj_f else None,
    )

    pj_d = data["penjualan_detail"]
    _write_sheet(
        wb, "Penjualan Detail Luar",
        PENJUALAN_DETAIL_COLS, pj_d,
        total_cols={8: sum(int(r.get("qty") or 0) for r in pj_d), 10: ""} if pj_d else None,
    )

    pb_f = data["pembelian_faktur"]
    _write_sheet(
        wb, "Pembelian Faktur",
        PEMBELIAN_FAKTUR_COLS, pb_f,
        total_cols={
            5: sum(int(r.get("total_line") or 0) for r in pb_f),
            6: sum(int(r.get("line_dalam_41k") or 0) for r in pb_f),
            7: sum(int(r.get("line_luar_41k") or 0) for r in pb_f),
            8: sum(int(r.get("qty_dalam_41k") or 0) for r in pb_f),
            9: sum(int(r.get("qty_luar_41k") or 0) for r in pb_f),
        } if pb_f else None,
    )

    pb_d = data["pembelian_detail"]
    _write_sheet(
        wb, "Pembelian Detail Luar",
        PEMBELIAN_DETAIL_COLS, pb_d,
        total_cols={6: sum(int(r.get("qty") or 0) for r in pb_d)} if pb_d else None,
    )

    br = data["barang_ringkasan"]
    _write_sheet(
        wb, "Barang Luar Ringkasan",
        BARANG_RINGKASAN_COLS, br,
        total_cols={
            4: sum(int(r.get("faktur_jual") or 0) for r in br),
            5: sum(int(r.get("qty_jual") or 0) for r in br),
            7: sum(int(r.get("faktur_beli") or 0) for r in br),
            8: sum(int(r.get("qty_beli") or 0) for r in br),
        } if br else None,
    )

    tr = data["transfer"]
    _write_sheet(
        wb, "Transfer Luar",
        TRANSFER_COLS, tr,
        total_cols={
            8: sum(int(r.get("jumlah_transfer") or 0) for r in tr),
            9: sum(int(r.get("jumlah_terima") or 0) for r in tr),
        } if tr else None,
    )

    wb.save(path)
